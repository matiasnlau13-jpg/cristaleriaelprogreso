"""Genera planilla_piezas.xlsx a partir de los datos transcriptos de las fotos del cuaderno.

Para agregar piezas, sumar dicts en DATOS y volver a correr:
    python planilla/generar_planilla.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SALIDA = Path(__file__).with_name("planilla_piezas.xlsx")

# Columnas de cada hoja. "P" = peso, "B" = base.
# "Cliente" = nombres que en el cuaderno están entre comillas o sueltos en lápiz.
COLUMNAS = {
    "VASOS": ["Nombre", "Cliente", "Altura", "Peso (P)", "Base (B)"],
    "COPAS": ["Nombre", "Cliente", "Pierna (estirada / pegada)", "Corte caliente c/c (altura)", "c/", "Peso (P)", "Base (B)", "Pie", "Pinza", "Molde"],
    "FLORERO": ["Nombre", "Cliente", "Capacidad", "Corte caliente c/c (altura)", "Otro corte", "Peso (P)", "Base (B)", "Nota en lápiz"],
    "TULIPA": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)"],
    "PAMELA": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)"],
    "BOTELLON": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)"],
    "TARRO": ["Nombre", "Cliente", "Capacidad", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)", "Nota en lápiz"],
    "CENTRO": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)"],
    "TAPA": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)", "Molde"],
    "JARRA": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)"],
    "ACEITERA": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)"],
    "OTROS": ["Nombre", "Cliente", "Corte caliente c/c (altura)", "Peso (P)", "Base (B)"],
    "PRENSA": ["Tipo", "Nombre", "Cliente", "Peso (P)", "Base (B)"],
}
EXTRA = ["Sección del cuaderno", "Foto Nº", "Observaciones"]

# Filas por hoja: cada pieza es un dict {columna: valor}; las columnas que faltan quedan vacías.
DATOS = {
    "COPAS": [
        {"Nombre": "Copa Romina Vino", "Pierna (estirada / pegada)": "Pegada", "Corte caliente c/c (altura)": "15,3", "Peso (P)": 200, "Base (B)": 230, "Pie": "Lorena V.B", "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Presidente Champ", "Peso (P)": 180, "Base (B)": 200, "Sección del cuaderno": "Copas", "Foto Nº": 2, "Cliente": "Renato"},
        {"Nombre": "Copa Clarito Inés De los Santos", "Peso (P)": 200, "Base (B)": 230, "Pie": "Presi V.T", "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Prince Agua", "Pierna (estirada / pegada)": "Pegada", "Corte caliente c/c (altura)": "19,7", "Peso (P)": 220, "Base (B)": 220, "Pie": "London 98", "Sección del cuaderno": "Copas", "Foto Nº": 2, "Observaciones": "El peso está remarcado/sobrescrito: confirmar que es 220"},
        {"Nombre": "Copa Barón B", "Pierna (estirada / pegada)": "Estirada", "Peso (P)": 250, "Base (B)": 200, "Pie": "London 98", "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Chandón 2021", "Pierna (estirada / pegada)": "Estirada", "Peso (P)": 230, "Base (B)": 200, "Pie": "London 98", "Sección del cuaderno": "Copas", "Foto Nº": 2, "Observaciones": "El pie no está escrito: hay comillas (\") debajo de \"Pie London 98\" de la fila de arriba"},
        {"Nombre": "Copa Magnífica", "Pierna (estirada / pegada)": "Estirada", "Peso (P)": 160, "Base (B)": 200, "Pie": "Delise", "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Venus Champ", "Pierna (estirada / pegada)": "Estirada", "Peso (P)": 160, "Base (B)": 220, "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Flavia Champ", "Pierna (estirada / pegada)": "Estirada", "Peso (P)": 160, "Base (B)": 220, "Sección del cuaderno": "Copas", "Foto Nº": 2, "Cliente": "Renato"},
        {"Nombre": "Copa Ceci Flauta", "Pierna (estirada / pegada)": "Estirada", "c/": "24", "Peso (P)": 170, "Base (B)": 200, "Pie": "Presi Agua", "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Him-Chan Agua", "c/": "15", "Peso (P)": 150, "Base (B)": 230, "Pie": "Venus V.B", "Sección del cuaderno": "Copas", "Foto Nº": 2, "Cliente": "Renato"},
        {"Nombre": "Copa Retro", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 180, "Base (B)": 220, "Pie": "Mod II Agua", "Sección del cuaderno": "Copas", "Foto Nº": 2, "Cliente": "Renato"},
        {"Nombre": "Copa Premium Mediana", "Peso (P)": 230, "Base (B)": 220, "Pie": "Mini Marisco", "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Premium Chica", "Peso (P)": 220, "Base (B)": 220, "Pie": "Mod II Agua", "Sección del cuaderno": "Copas", "Foto Nº": 2},
        {"Nombre": "Copa Silvia Agua", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 180, "Base (B)": 220, "Pie": "Silvia Agua", "Sección del cuaderno": "Copas", "Foto Nº": 2, "Cliente": "Renato"},
        {"Nombre": "Copa Expresión", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 230, "Base (B)": 200, "Pie": "Premium", "Sección del cuaderno": "Copas", "Foto Nº": 3},
        {"Nombre": "Copa Martini Champ", "Pierna (estirada / pegada)": "Alta", "Peso (P)": 260, "Base (B)": 200, "Pie": "Eva Champ", "Sección del cuaderno": "Copas", "Foto Nº": 3, "Cliente": "C. Ruiz"},
        {"Nombre": "Copa Nueva Progreso", "Corte caliente c/c (altura)": "21,0", "Peso (P)": 300, "Base (B)": 200, "Pie": "Premium", "Pinza": "19", "Sección del cuaderno": "Copas", "Foto Nº": 3},
        {"Nombre": "Copa Hurricane", "Corte caliente c/c (altura)": "20,5", "Peso (P)": 300, "Base (B)": 200, "Pie": "Coñac London 98", "Pinza": "90", "Sección del cuaderno": "Copas", "Foto Nº": 3, "Observaciones": "Pinza: se lee 90 (las demás dicen 19). Confirmar"},
        {"Nombre": "Copa Patagonia Cerveza", "Corte caliente c/c (altura)": "15,3", "Peso (P)": 250, "Base (B)": 230, "Pie": "London 98", "Pinza": "19", "Sección del cuaderno": "Copas", "Foto Nº": 3},
        {"Nombre": "Copa Taby Cerveza", "Corte caliente c/c (altura)": "14,3", "Peso (P)": 280, "Base (B)": 200, "Pie": "Venus Agua", "Pinza": "19", "Sección del cuaderno": "Copas", "Foto Nº": 3},
        {"Nombre": "Copa Ani", "Pierna (estirada / pegada)": "Pegada (pierna fina)", "Peso (P)": 280, "Base (B)": 200, "Pie": "Marisco Deisy", "Molde": "Copa \"G\"", "Sección del cuaderno": "Copas", "Foto Nº": 3, "Observaciones": "Molde copa \"G\": puede ser G o 6. Confirmar En el cuaderno dice Destino: \"German\".", "Cliente": "German"},
        {"Nombre": "Copa Checo Flauta", "Pierna (estirada / pegada)": "Estirada", "c/": "22,0", "Peso (P)": 170, "Base (B)": 230, "Pie": "Presi Agua", "Molde": "Erika de Rosemblet", "Sección del cuaderno": "Copas", "Foto Nº": 3, "Cliente": "Renato"},
        {"Nombre": "Copa Premium Borgoña", "Peso (P)": 280, "Base (B)": 200, "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa Premium Borgoña R", "Peso (P)": 280, "Base (B)": 200, "Sección del cuaderno": "Copas", "Foto Nº": 7, "Observaciones": "Renglón con comillas de repetición debajo de Premium Borgoña, con \"R\" y \"C. Ruiz\". ¿Es otra versión de la misma copa?", "Cliente": "C. Ruiz"},
        {"Nombre": "Copa G", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 260, "Base (B)": 230, "Pie": "Andrea Agua", "Sección del cuaderno": "Copas", "Foto Nº": 7, "Observaciones": "¿Es la misma copa \"G\" que figura como molde de la Copa Ani (foto 3)?"},
        {"Nombre": "Copa Beefeater", "Pierna (estirada / pegada)": "Alta", "Peso (P)": 280, "Base (B)": 200, "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa Bordeau Agua", "Peso (P)": 280, "Base (B)": 200, "Pie": "7x7 Champ", "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa H", "Cliente": "Renato", "Pierna (estirada / pegada)": "Alta", "Peso (P)": 200, "Base (B)": 220, "Pie": "Mod II Agua", "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa Cabernet", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 300, "Base (B)": 180, "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa Terraza", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 280, "Base (B)": 200, "Pie": "Cheval", "Molde": "Cheval", "Sección del cuaderno": "Copas", "Foto Nº": 7, "Observaciones": "La pierna tiene comillas de repetición (\") debajo de \"P.Peg\" de la Cabernet. \"Molde Cheval\" está en lápiz"},
        {"Nombre": "Copa 2 Tiempo", "Peso (P)": 280, "Base (B)": 200, "Pie": "Clásica", "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa Alma", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 280, "Base (B)": 200, "Sección del cuaderno": "Copas", "Foto Nº": 7, "Observaciones": "Dice \"Pie\" pero no está escrito cuál"},
        {"Nombre": "Copa Degustación", "Peso (P)": 150, "Base (B)": 240, "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa M-H Vino", "Cliente": "C. Ruiz", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 200, "Base (B)": 200, "Pie": "London 98", "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa Vitalia", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 280, "Base (B)": 200, "Pie": "Clásica", "Sección del cuaderno": "Copas", "Foto Nº": 7, "Cliente": "Renato"},
        {"Nombre": "Copa Vinarte", "Pierna (estirada / pegada)": "Alta y Pegada", "Peso (P)": 160, "Base (B)": 200, "Pie": "Presi Vino Tinto", "Sección del cuaderno": "Copas", "Foto Nº": 7, "Observaciones": "Dice \"P.Alta\" y \"P.Peg\" a la vez", "Cliente": "Renato"},
        {"Nombre": "Copa Soire Flauta", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 200, "Base (B)": 220, "Pie": "Presi Agua", "Sección del cuaderno": "Copas", "Foto Nº": 7},
        {"Nombre": "Copa Brunelo Vino", "Pierna (estirada / pegada)": "Pegada", "Peso (P)": 250, "Base (B)": 200, "Pie": "Brenda Agua", "Sección del cuaderno": "Copas", "Foto Nº": 7},
    ],
    "FLORERO": [
        {"Nombre": "Florero Cónico", "Capacidad": "1 L", "Corte caliente c/c (altura)": "14,5", "Peso (P)": 900, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Florero Cónico", "Capacidad": "2 L", "Corte caliente c/c (altura)": "18,0", "Peso (P)": 1000, "Base (B)": 100, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Florero 12x30", "Corte caliente c/c (altura)": "30,0", "Peso (P)": 1000, "Base (B)": 100, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Florero 17x50", "Corte caliente c/c (altura)": "47,0", "Peso (P)": 2500, "Base (B)": 100, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Florero 17x50 2 cortes", "Otro corte": "2 cortes", "Peso (P)": 3000, "Base (B)": 100, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Florero Lila", "Corte caliente c/c (altura)": "15,3", "Peso (P)": 650, "Base (B)": 200, "Nota en lápiz": "Chico", "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Florero Lila Diagonal", "Corte caliente c/c (altura)": "Diagonal", "Peso (P)": 850, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 4, "Observaciones": "Dice \"c/c Diagonal\" sin altura: ¿corte en diagonal?"},
        {"Nombre": "Florero Potiche", "Corte caliente c/c (altura)": "23,5", "Peso (P)": 1200, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Florero Cuadrado Grande", "Peso (P)": 1200, "Base (B)": 200, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Florero Mod 11", "Corte caliente c/c (altura)": "26,5", "Otro corte": "c/ 24 cm", "Peso (P)": 700, "Base (B)": 200, "Sección del cuaderno": "Sin título", "Foto Nº": 5, "Observaciones": "Abajo dice: corte c/24cm \"Cenatiempo\". Confirmar cómo se escribe Cenatiempo", "Cliente": "Cenatiempo"},
        {"Nombre": "Florero Cónico Alto", "Corte caliente c/c (altura)": "28,0", "Peso (P)": 550, "Base (B)": 250, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Florero Cónico Bajo", "Corte caliente c/c (altura)": "19,3", "Peso (P)": 400, "Base (B)": 250, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
    ],
    "TULIPA": [
        {"Nombre": "Tulipa Mini Mum", "Peso (P)": 900, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Tulipa Rochester", "Corte caliente c/c (altura)": "20,5", "Peso (P)": 1500, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
    ],
    "PAMELA": [
        {"Nombre": "Pamela x21", "Corte caliente c/c (altura)": "9,0", "Peso (P)": 1300, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Pamela x16", "Corte caliente c/c (altura)": "7,6", "Peso (P)": 700, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Pamela x11", "Peso (P)": 400, "Base (B)": 300, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Pamela x10", "Corte caliente c/c (altura)": "5,3", "Peso (P)": 350, "Base (B)": 250, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
    ],
    "BOTELLON": [
        {"Nombre": "Botellón Neo", "Peso (P)": 300, "Base (B)": 300, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Botellón Inti Chico", "Peso (P)": 400, "Base (B)": 200, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Botellón Mini", "Peso (P)": 500, "Base (B)": 250, "Sección del cuaderno": "Sin título", "Foto Nº": 5},
        {"Nombre": "Botellón Decanter", "Peso (P)": 1300, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
        {"Nombre": "Botellón Decantador Progreso", "Peso (P)": 1300, "Base (B)": 150, "Sección del cuaderno": "Sin título", "Foto Nº": 6, "Observaciones": "Escrito en lápiz debajo del Botellón Decanter. ¿Es la misma pieza con otro nombre?"},
        {"Nombre": "Botellón Nahuel", "Peso (P)": 1300, "Base (B)": 100, "Sección del cuaderno": "Sin título", "Foto Nº": 6, "Observaciones": "Escrito en lápiz"},
    ],
    "TARRO": [
        {"Nombre": "Tarro Cónico", "Capacidad": "1/2 L", "Corte caliente c/c (altura)": "13,0", "Peso (P)": 500, "Base (B)": 250, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Tarro Claret", "Capacidad": "1 1/2 L", "Corte caliente c/c (altura)": "20,1", "Peso (P)": 1100, "Base (B)": 120, "Nota en lápiz": "B:55", "Sección del cuaderno": "Sin título", "Foto Nº": 4, "Observaciones": "Tiene un segundo \"B:55\" al final del renglón. ¿Qué es?"},
        {"Nombre": "Tarro Bocha", "Capacidad": "3/4 L", "Corte caliente c/c (altura)": "11,0", "Peso (P)": 300, "Base (B)": 250, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Tarro M-13x8,5", "Corte caliente c/c (altura)": "8,5", "Peso (P)": 800, "Base (B)": 200, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Tarro M-13x15", "Corte caliente c/c (altura)": "15,0", "Peso (P)": 1000, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Tarro M-13x20", "Corte caliente c/c (altura)": "20,0", "Peso (P)": 1000, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Tarro M-13x25", "Corte caliente c/c (altura)": "25,0", "Peso (P)": 1100, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
        {"Nombre": "Tarro M-13x30", "Corte caliente c/c (altura)": "30,0", "Peso (P)": 1200, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 4},
    ],
    "CENTRO": [
        {"Nombre": "Centro Oreford Grande", "Peso (P)": 2800, "Base (B)": 100, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
        {"Nombre": "Centro Porta Vela", "Corte caliente c/c (altura)": "9,5", "Peso (P)": 1300, "Base (B)": 100, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
    ],
    "TAPA": [
        {"Nombre": "Tapa Quesera 1/2 Esfera", "Corte caliente c/c (altura)": "c/c (sin altura)", "Peso (P)": 1300, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 6, "Observaciones": "Dice c/c pero no tiene la altura"},
        {"Nombre": "Tapa Quesera con Botón", "Peso (P)": 3000, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
        {"Nombre": "Tapa 17x20", "Corte caliente c/c (altura)": "20,0", "Peso (P)": 1200, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
        {"Nombre": "Tapa 17x13", "Corte caliente c/c (altura)": "13,0", "Peso (P)": 1200, "Base (B)": 120, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
        {"Nombre": "Tapa M-13", "Peso (P)": 400, "Base (B)": 250, "Molde": "Also", "Sección del cuaderno": "Sin título", "Foto Nº": 6, "Observaciones": "Molde: se lee \"Also\". Confirmar"},
    ],
    "JARRA": [
        {"Nombre": "Jarra Carafe con Pico y Asa", "Corte caliente c/c (altura)": "16,0", "Peso (P)": 550, "Base (B)": 55, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
    ],
    "ACEITERA": [
        {"Nombre": "Aceitera Cónico Chico con Pico", "Peso (P)": 400, "Base (B)": 250, "Sección del cuaderno": "Sin título", "Foto Nº": 6},
    ],
    "OTROS": [
        {"Nombre": "Bombe Gigante", "Corte caliente c/c (altura)": "15,8", "Peso (P)": 700, "Base (B)": 200, "Sección del cuaderno": "Sin título", "Foto Nº": 6, "Observaciones": "No entra en ninguna categoría. ¿Dónde va?"},
    ],
    "PRENSA": [
        {"Tipo": "Plato", "Nombre": "Plato Pizza", "Peso (P)": 300, "Base (B)": 500, "Sección del cuaderno": "Prensa", "Foto Nº": 1},
        {"Tipo": "Copa", "Nombre": "Copa Alaska", "Peso (P)": 320, "Base (B)": 450, "Sección del cuaderno": "Prensa", "Foto Nº": 1, "Observaciones": "En lápiz: \"Requem. 8\" y un \"300\" debajo del B:450 (¿B corregido a 300?)"},
        {"Tipo": "Faro", "Nombre": "Faro Puma", "Peso (P)": 130, "Base (B)": 450, "Sección del cuaderno": "Prensa", "Foto Nº": 1},
        {"Tipo": "Faro", "Nombre": "Faro Ford", "Peso (P)": 452, "Base (B)": 600, "Sección del cuaderno": "Prensa", "Foto Nº": 1},
        {"Tipo": "Copa", "Nombre": "Copa Milk Shake", "Peso (P)": 570, "Base (B)": 250, "Sección del cuaderno": "Prensa", "Foto Nº": 1},
        {"Tipo": "Filtro", "Nombre": "Filtro Nº 2", "Peso (P)": 245, "Base (B)": 500, "Sección del cuaderno": "Prensa", "Foto Nº": 1},
        {"Tipo": "Filtro", "Nombre": "Filtro Nº 4", "Peso (P)": 40, "Base (B)": 600, "Sección del cuaderno": "Prensa", "Foto Nº": 1, "Observaciones": "Escrito en lápiz; \"SINAC\" con flecha desde B:600", "Cliente": "SINAC"},
        {"Tipo": "Base Velador", "Nombre": "Base Velador Grande", "Peso (P)": 465, "Base (B)": 400, "Sección del cuaderno": "Prensa", "Foto Nº": 1},
        {"Tipo": "Base Velador", "Nombre": "Base Velador Chico", "Peso (P)": 280, "Base (B)": 600, "Sección del cuaderno": "Prensa", "Foto Nº": 1, "Observaciones": "\"Rococo\" está escrito debajo del Chico", "Cliente": "Rococo"},
        {"Tipo": "Lupa", "Nombre": "Lupa Nº 3", "Peso (P)": 88, "Base (B)": 500, "Sección del cuaderno": "Prensa", "Foto Nº": 1, "Cliente": "FYW"},
        {"Tipo": "Lupa", "Nombre": "Lupa Nº 4", "Peso (P)": 350, "Base (B)": 500, "Sección del cuaderno": "Prensa", "Foto Nº": 1},
        {"Tipo": "Vidrio", "Nombre": "Vidrio U", "Peso (P)": 620, "Base (B)": 300, "Sección del cuaderno": "Prensa", "Foto Nº": 1, "Observaciones": "Escrito en lápiz"},
        {"Tipo": "Baliza", "Nombre": "Baliza Anillada Grande", "Peso (P)": 650, "Base (B)": 300, "Sección del cuaderno": "Prensa", "Foto Nº": 1, "Cliente": "BALIZAR"},
    ],
}

FUENTE = "Arial"
ENCABEZADO_FILL = PatternFill("solid", start_color="1F4E78")
DUDA_FILL = PatternFill("solid", start_color="FFF2CC")
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for hoja, cols in COLUMNAS.items():
        ws = wb.create_sheet(hoja)
        encabezados = cols + EXTRA
        ws.append(encabezados)
        for celda in ws[1]:
            celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
            celda.fill = ENCABEZADO_FILL
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.border = BORDE
        for fila in DATOS.get(hoja, []):
            ws.append([fila.get(c, "") for c in encabezados])
            r = ws.max_row
            for celda in ws[r]:
                celda.font = Font(name=FUENTE)
                celda.border = BORDE
                celda.alignment = Alignment(vertical="center", wrap_text=True)
            if fila.get("Observaciones"):
                for celda in ws[r]:
                    celda.fill = DUDA_FILL
        for i, titulo in enumerate(encabezados, start=1):
            letra = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letra].width = 55 if titulo == "Observaciones" else max(14, len(titulo) + 4)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
    wb.save(SALIDA)
    print(f"Guardado {SALIDA}")


if __name__ == "__main__":
    main()
